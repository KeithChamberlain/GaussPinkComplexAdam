"""Analysis of Initializer and Optimizer Combinations."""
import sys
import json
import numpy as np
import pandas as pd
import openpyxl
from plotly.subplots import make_subplots
import plotly.graph_objects as go
import statsmodels.formula.api as smf
from scipy.stats import t, sem


def to_data(df, gb, keep, keys):
    results = dict()
    for key in keys:
        results[key] = df.groupby(gb)[keep[key]].agg("mean")
    return results


def em(codes, start=0, coefs=None, transpose=False):
    if transpose:
        return np.transpose(codes.iloc[:, start:].to_numpy()).dot(coefs['Coef.'])
    return codes.iloc[:, start:].to_numpy().dot(coefs['Coef.']) 


if __name__ == "__main__":
    workbook = str(sys.argv[1])
    output_workbook = str(sys.argv[2])
    confidence = float(sys.argv[3])
    # Prelim Declarations
    formulas = {
        "loss": "val_loss ~ gauss_vs_pink + adam_vs_complex + interaction",
        "acc": "val_acc ~ gauss_vs_pink + adam_vs_complex + interaction",
        "train": "train ~ gauss_vs_pink + adam_vs_complex + interaction"
    }
    keys = ['loss', 'acc', 'train']
    keep = {
        "loss": ["val_loss", "gauss_vs_pink", "adam_vs_complex", "interaction"],
        "acc": ["val_acc", "gauss_vs_pink", "adam_vs_complex", "interaction"],
        "train": ["train", "gauss_vs_pink", "adam_vs_complex", "interaction"],
    }

    # Read Workbook
    wkbook = openpyxl.load_workbook(workbook, data_only=True)
    df_list = list()
    for sheet in wkbook.sheetnames:
        print(sheet)
        if sheet == "Sheet1":
            pass
        else:
            rows = list(wkbook[sheet].iter_rows(values_only=True))
            header = rows[0]
            data = rows[1:]
            df = pd.DataFrame(data, columns=header)
            df['sheet'] = sheet
            df_list.append(df)
    wkbook.close()
    workbook_as_df = pd.concat(df_list)

    # Make Grouping Column
    workbook_as_df['group'] = workbook_as_df['init'] + ", " + workbook_as_df['opt'].str.replace("torch_", "").replace("complex_adam", "complex")
    gb = workbook_as_df.groupby(["group"])[["val_loss", "val_acc", "train"]].agg(["mean", sem])
    print(gb)
    print(workbook_as_df.head(10))


    # Set Contrasts
    #       gauss_vs_pink. adam-vs_complex.           int
    # gauss, adam      -.5            -.5           -.25
    # gauss, complex   -.5             .5            .25
    # pink, adam       .5             -.5            .25
    # pink, complex    .5              .5           -.25
    codes = pd.DataFrame({
        "means": workbook_as_df.loc[:, "group"].unique(),
        "intercept": np.repeat(1, 4),
        "gauss_vs_pink": np.array([-.5, -.5, +.5, +.5]),
        "adam_vs_complex": np.array([-.5, +.5, -.5, +.5]),
        "interaction": np.array([-.25, +.25, +.25, -.25]),
    })
    print(codes) # Check the order


    workbook_as_df['gauss_vs_pink'] = (
        -.5*(workbook_as_df['group'] == "gauss, adam").astype('int') -
        .5*(workbook_as_df['group'] == "gauss, complex").astype('int') +
        .5*(workbook_as_df['group'] == "pink, adam").astype('int') +
        .5*(workbook_as_df['group'] == "pink, complex").astype('int')
    )
    workbook_as_df['adam_vs_complex'] = (
        -.5*(workbook_as_df['group'] == "gauss, adam").astype('int') +
        .5*(workbook_as_df['group'] == "gauss, complex").astype('int') -
        .5*(workbook_as_df['group'] == "pink, adam").astype('int') +
        .5*(workbook_as_df['group'] == "pink, complex").astype('int')
    )
    workbook_as_df['interaction'] = (
        -.25*(workbook_as_df['group'] == "gauss, adam").astype('int') +
        .25*(workbook_as_df['group'] == "gauss, complex").astype('int') +
        .25*(workbook_as_df['group'] == "pink, adam").astype('int') -
        .25*(workbook_as_df['group'] == "pink, complex").astype('int')
    )
    check1 = workbook_as_df.loc[workbook_as_df["sheet"] == "gauss_torch_adam", ['seed']]
    check2 = workbook_as_df.loc[workbook_as_df["sheet"] == "gauss_complex_adam", ['seed']]
    check3 = workbook_as_df.loc[workbook_as_df["sheet"] == "pink_torch_adam", ['seed']]
    check4 = workbook_as_df.loc[workbook_as_df["sheet"] == "pink_complex_adam", ['seed']]
    print(any(check1['seed'] != check2['seed']))
    print(any(check1['seed'] != check3['seed']))
    print(any(check1['seed'] != check4['seed']))
    print(any(check2['seed'] != check3['seed']))
    print(any(check2['seed'] != check4['seed']))
    print(any(check3['seed'] != check4['seed']))
    
    fits = {
        key: smf.ols(formula=formulas[key], data=workbook_as_df).fit()
        for key in keys
    }
    print(fits['loss'])
    # Review Original Fits
    summaries = {
        key: fits[key].summary2().tables[1]
        for key in keys
    }
    print(summaries)
    # Get values for CI for each group.
    print("groups")
    print()
    margins = dict(
        loss = em(codes, 1, summaries['loss']),
        acc = em(codes, 1, summaries['acc']),
        train = em(codes, 1, summaries['train']),
    )
    print(margins)
    print("df_resid")
    print(fits['loss'].df_resid)
    cis = {
        key: t.interval(1-(1-confidence)/2, df=fits[key].df_resid, loc=margins[key], scale=gb[(keep[key][0], "sem")])
        for key in keys
    }
    print(list(cis))


    figx = make_subplots(rows=1, cols=3,
        x_title = "Mean Differences",
        y_title = "Scores",
        specs = [[{"secondary_y": False}, {"secondary_y": False}, {"secondary_y": False}]],
        subplot_titles = ["Loss", "Accuracy", "Training Time (s)"]
    )
    fig_orig = make_subplots(rows=2, cols=1,
        x_title = "Group Means",
        y_title = "Scores",
        specs=[[{"secondary_y": True}], [{"secondary_y":False}]],
        subplot_titles = ["Loss and accuracy metrics", "Time metric"]
    )
    i=1
    for key, label, dv in zip(keys, ['Loss (left)', 'Accuracy (middle)', 'Training time (right)'], ['val_loss', 'val_acc', 'train']):
        index = summaries[key].index.values[1:]
        print(index)
        figx.add_trace(go.Bar(
            name=label,
            x=index,
            y=summaries[key].loc[index, 'Coef.'],
            error_y=dict(
                type="data",
                symmetric=False,
                array=summaries[key].loc[index, '0.975]'] - summaries[key].loc[index, 'Coef.'],
                arrayminus=summaries[key].loc[index,'Coef.'] - summaries[key].loc[index, '[0.025']
            )
        ), secondary_y=False, row=1, col=i)
        i+=1
        fig_orig.add_trace(go.Scatter(
            name=label,
            x=index,
            y=gb[(dv, 'mean')], 
            mode="markers",
            error_y=dict(
                type='data',
                symmetric=True,
                array=summaries[key]['Std.Err.']
            )
        ), secondary_y=False if i != 2 else True, row=1 if i != 3 else 2, col=1)

    figx.update_layout(
        title=dict(
            text="Mean differences for transformer neural network metrics for initializer and optimizer types",
            subtitle=dict(text="With 95% confidence interval")))

    fig_orig.update_yaxes(title_text = "Loss", row=1, col=1, secondary_y=False)
    fig_orig.update_yaxes(title_text = "Accuracy", row=1, col=1, secondary_y=True)
    fig_orig.update_yaxes(title_text = "Training Time (s)", row=2, col=1, secondary_y=False)
    fig_orig.update_layout(
        title=dict(
            text="Means for transformer neural network metrics for initializer and optimizer types",
            subtitle=dict(text="With standard error bars")))

    with open("combined_dashboard.html", "w") as f:
        f.write("<html><head><title>Dashboard</title></head><body>")
        f.write(figx.to_html(full_html=False, include_plotlyjs='cdn'))
        f.write(fig_orig.to_html(full_html=False, include_plotlyjs='cdn'))
        f.write("</body></html>")

    summaries['loss'].to_excel(output_workbook, sheet_name="LOSS_OLS")

    with pd.ExcelWriter(output_workbook, engine="openpyxl", mode="a") as writer:
        gb.to_excel(writer, sheet_name="MEANS")
        for key in keys:
            if key != 'loss':
                summaries[key].to_excel(writer, sheet_name=key.upper()+"_OLS")
